"""Real in-memory document fixtures; no mocks of extraction or prediction."""

from io import BytesIO
import zipfile
import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from starlette.requests import Request
from api.main import create_app
from api.uploads import read_cv_upload
from src.extract import MAX_FILE_BYTES, FileValidationError, NO_TEXT

TEXT = "Accountant managing audits, financial statements and tax reporting"


def pdf_bytes(text=TEXT):
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 30 750 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = stream
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_bytes(text=TEXT, table=False):
    document = Document()
    if table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = text
    else:
        document.add_paragraph(text)
    out = BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.fixture(scope="module")
def client():
    with TestClient(create_app()) as client:
        yield client


@pytest.mark.parametrize(
    "extension,make",
    [("pdf", pdf_bytes), ("docx", docx_bytes), ("txt", lambda: TEXT.encode())],
)
def test_valid_file_same_pipeline(client, extension, make):
    response = client.post("/predict/file", files={"file": (f"cv.{extension}", make())})
    assert response.status_code == 200, response.text
    result = response.json()
    expected = client.post("/predict", json={"text": TEXT}).json()
    assert result == {**expected, "source": "file", "text_preview": TEXT}


@pytest.mark.parametrize(
    "filename,content,status",
    [
        ("cv.exe", TEXT.encode(), 415),
        ("cv.pdf", TEXT.encode(), 415),
        ("cv.docx", b"PK\x03\x04not a zip", 415),
        ("cv.txt", b"\x00\x01binary cv content", 415),
        ("cv.txt", pdf_bytes(), 415),
        ("cv.pdf", b"%PDFbroken", 422),
        ("cv.txt", b"a" * (MAX_FILE_BYTES + 1), 413),
        ("cv.txt", b"", 422),
        ("cv.docx", b"", 422),
        ("cv.pdf", b"", 422),
        ("cv.txt", b"hi", 422),
    ],
    ids=[
        "wrong-extension",
        "fake-pdf",
        "fake-docx",
        "binary-txt",
        "renamed-pdf",
        "broken-pdf",
        "oversized",
        "empty-txt",
        "empty-docx",
        "empty-pdf",
        "too-short",
    ],
)
def test_reject_invalid_upload(client, filename, content, status):
    response = client.post("/predict/file", files={"file": (filename, content)})
    assert response.status_code == status
    assert isinstance(response.json()["detail"], str)


def test_image_only_pdf(client):
    response = client.post("/predict/file", files={"file": ("scan.pdf", pdf_bytes(""))})
    assert response.status_code == 422
    assert response.json() == {"detail": NO_TEXT}


def test_docx_requires_document_xml(client):
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("unrelated.txt", TEXT)
    assert (
        client.post(
            "/predict/file", files={"file": ("cv.docx", out.getvalue())}
        ).status_code
        == 415
    )


def test_docx_table_extraction(client):
    response = client.post(
        "/predict/file", files={"file": ("cv.docx", docx_bytes(table=True))}
    )
    assert response.status_code == 200
    assert response.json()["text_preview"] == TEXT


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "cp1252"])
def test_text_encodings(client, encoding):
    text = TEXT + " café résumé"
    response = client.post(
        "/predict/file", files={"file": ("CV.TXT", text.encode(encoding))}
    )
    assert response.status_code == 200
    assert response.json()["text_preview"] == text


def test_no_disk_spooling(client, monkeypatch):
    def fail(*args, **kwargs):
        raise AssertionError("Uploads must not use temporary disk files")

    monkeypatch.setattr("starlette.formparsers.SpooledTemporaryFile", fail)
    monkeypatch.setattr("tempfile.TemporaryFile", fail)
    content = BytesIO(docx_bytes())
    with zipfile.ZipFile(content, "a", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("padding.bin", b"x" * (2 * 1024 * 1024))
    assert (
        client.post(
            "/predict/file", files={"file": ("cv.docx", content.getvalue())}
        ).status_code
        == 200
    )


def test_multiple_files_rejected(client):
    response = client.post(
        "/predict/file",
        files=[("file", ("one.txt", TEXT)), ("file", ("two.txt", TEXT))],
    )
    assert response.status_code == 400


def test_wrong_field(client):
    assert (
        client.post("/predict/file", files={"resume": ("cv.txt", TEXT)}).status_code
        == 400
    )


def test_missing_multipart(client):
    assert client.post("/predict/file", json={"text": TEXT}).status_code == 400


def test_streamed_limit_without_content_length():
    import asyncio

    received = 0

    async def receive():
        nonlocal received
        received += 1
        if received == 1:
            return {
                "type": "http.request",
                "body": b'--cv\r\nContent-Disposition: form-data; name="file"; filename="cv.txt"\r\n\r\n',
                "more_body": True,
            }
        return {"type": "http.request", "body": b"x" * 65536, "more_body": True}

    request = Request(
        {
            "type": "http",
            "headers": [(b"content-type", b"multipart/form-data; boundary=cv")],
        },
        receive,
    )
    with pytest.raises(FileValidationError) as error:
        asyncio.run(read_cv_upload(request))
    assert error.value.status_code == 413
    assert received <= 82  # Stops consumption; never waits for an unbounded body.


def test_incomplete_multipart(client):
    body = (
        b'--cv\r\nContent-Disposition: form-data; name="file"; filename="cv.txt"\r\n\r\n'
        + TEXT.encode()
    )
    assert (
        client.post(
            "/predict/file",
            content=body,
            headers={"Content-Type": "multipart/form-data; boundary=cv"},
        ).status_code
        == 400
    )


def test_preview_bounded_and_short_flag(client):
    text = ("accounting financial audit reporting " * 30).strip()
    result = client.post("/predict/file", files={"file": ("cv.txt", text)}).json()
    assert result["text_preview"] == text[:500]
    assert result["text_stats"] == {"word_count": 120, "short_input": False}


def test_top_predictions_and_terms_are_from_artifacts(client):
    predictor = client.app.state.predictor
    result = client.post("/predict", json={"text": TEXT}).json()
    vector = predictor.vectorizer.transform([TEXT])
    probabilities = predictor.model.predict_proba(vector)[0]
    order = np.argsort(-probabilities, kind="stable")[:3]
    assert result["top_predictions"] == [
        {
            "category": str(predictor.model.classes_[i]),
            "probability": float(probabilities[i]),
        }
        for i in order
    ]
    names = predictor.vectorizer.get_feature_names_out()
    scored = sorted(
        [
            (str(names[i]), value * predictor.model.feature_importances_[i])
            for i, value in zip(vector.indices, vector.data)
        ],
        key=lambda pair: -pair[1],
    )
    assert result["top_terms"] == [term for term, score in scored if score > 0][:8]
    assert result["text_stats"] == {"word_count": 8, "short_input": True}
    assert result["source"] == "text" and result["text_preview"] is None
